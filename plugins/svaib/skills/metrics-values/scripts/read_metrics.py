#!/usr/bin/env python3
"""read_metrics.py — чтение значений метрик из книги клиента по карте адресов.

Карта — эфемерный JSON, который агент выписывает, читая описание метрики в базе. В базе клиента
машинных артефактов не заводим: описание метрики остаётся единственным источником правды.

    python3 read_metrics.py --book <снимок>.xlsx --card <карта>.json [--upto 2026-09] [--json]

Что код выводит из книги сам: строку метрики (по меткам), ось периодов (шапка),
режим процента (формат ячейки). Книга, лист, метки и единица берутся из паспорта.
Масштаб по умолчанию 1; иной масштаб агент переносит из «Числа в книге».
labels — адрес для одного прогона, отдельного поля пути меток в паспорте нет.

Лист и метки всегда сопоставляются дословно, без нормализации и поиска подстрокой.
plan_source в записи метрики задаёт собственные book (локальный xlsx, необязательно),
sheet, labels, unit, scale и year плана. Его ряд присоединяется к plan этой метрики.
Путь book разрешается от каталога запуска, как --book; скрипт читает существующий локальный XLSX.

Сдвиг строки допустим, изменение имени листа или метки — отказ без числа.
При нескольких кандидатах или непонятной оси периодов число также не выдаётся.
Параметр strict_labels старых карт больше не ослабляет проверку.
"""
from __future__ import annotations

import argparse, datetime as dt, json, math, re, sys, zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("нужен openpyxl: pip install openpyxl", file=sys.stderr); raise SystemExit(2)

MONTHS = {'jan':1,'january':1,'янв':1,'январь':1,'feb':2,'february':2,'фев':2,'февраль':2,
          'mar':3,'march':3,'мар':3,'март':3,'apr':4,'april':4,'апр':4,'апрель':4,
          'may':5,'май':5,'мая':5,'jun':6,'june':6,'июн':6,'июнь':6,'jul':7,'july':7,'июл':7,'июль':7,
          'aug':8,'august':8,'авг':8,'август':8,'sep':9,'sept':9,'september':9,'сен':9,'сентябрь':9,
          'oct':10,'october':10,'окт':10,'октябрь':10,'nov':11,'november':11,'ноя':11,'ноябрь':11,
          'dec':12,'december':12,'дек':12,'декабрь':12}


def norm(s):
    return '' if s is None else re.sub(r'\s+', ' ', str(s)).strip().lower()


def is_percent_format(fmt):
    """Знак процента в числовом формате; проценты в строках/экранировании — литералы."""
    numeric = re.sub(r'"(?:[^"]|"")*"|\\.|_.|\*.', '', str(fmt))
    return '%' in numeric


# Шапка периодов бывает и машинной: «2026-09» (ISO), «2026/09», «09.2026» —
# так ведут книги, которые заполняет скрипт, а не человек. Год берётся из
# карты (`year`), как и для шапки из названий месяцев. Только год и месяц:
# «2026-09-01» и настоящие даты ячеек намеренно не месяц — дневной лист иначе
# читался бы как месячный, а это хуже отказа.
ISO_MONTH = re.compile(r'^(?:(\d{4})[-/.](\d{1,2})|(\d{1,2})[./](\d{4}))$')


def parse_month(v):
    k = norm(v)
    if not k:
        return None
    iso = ISO_MONTH.match(norm(v))
    if iso:
        month = int(iso.group(2) or iso.group(3))
        return month if 1 <= month <= 12 else None
    # Разбираем целую подпись месяца, а не начало слова: «маржа» не «март»,
    # «Jan plan» не факт за январь, «Jan-25» не январь неизвестного года.
    named = re.fullmatch(r'([a-zа-яё]+)\.?(?:\s+(?:19|20)\d{2})?', k)
    return MONTHS.get(named.group(1)) if named else None


def sheet_grid(ws):
    """Значения листа с forward-fill по объединённым ячейкам: метка группы стоит один раз."""
    grid = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[r - 1][c - 1] = v
    return grid


def read_numeric_cell(ws, formula_wb, sheet, coordinate, item, percent=None):
    """Read a stored numeric result; keep absence and errors explicit."""
    cell = ws[coordinate]
    formula_cell = (formula_wb[sheet][coordinate]
                    if formula_wb is not None else cell)
    value = cell.value
    if cell.data_type == "e" or (isinstance(value, str) and value.startswith("#")):
        return None, f"ошибка в ячейке источника ({value})"
    if value is None:
        reason = ("формула не имеет сохранённого результата"
                  if formula_cell.data_type == "f" else "ячейка пуста")
        return None, reason
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None, "ячейка содержит нечисловое значение"
    is_percentage = is_percent_format(cell.number_format) if percent is None else percent
    if is_percentage:
        return round(value * 100, 2), None
    return value * (item.get("scale") or 1), None


def find_period_axis(grid, max_scan=10):
    """Единственная шапка с минимум тремя месяцами; неоднозначную не выбираем."""
    axes = []
    for r in range(min(max_scan, len(grid))):
        cols = {}
        for c, v in enumerate(grid[r]):
            m = parse_month(v)
            if m and m not in cols:
                cols[m] = c
        if len(cols) >= 3:
            axes.append((r, cols))
    if len(axes) > 1:
        raise ValueError("несколько осей периодов: нужен проверенный читатель этой раскладки")
    best = axes[0] if axes else None
    if best:
        r, cols = best
        # Не выбираем первую колонку месяца: это может быть план или другой год.
        roles = {"план", "факт", "plan", "fact", "actual", "actuals", "budget", "%", "отклонение"}
        if any(norm(value) in roles for header in grid[:r] + grid[r + 1:r + 3]
               for value in header[min(cols.values()):max(cols.values()) + 1]):
            raise ValueError("роль колонки план/факт не определена: нужен проверенный читатель этой раскладки")
        months = [parse_month(value) for value in grid[r] if parse_month(value)]
        if len(months) != len(set(months)):
            raise ValueError("месяцы повторяются: годовой блок или роль колонки не определены")
    return best


def _match(path, want):
    i = 0
    for cell in path:
        if i < len(want) and want[i] and cell == want[i]:
            i += 1
    return i == len(want)


def literal(value):
    return '' if value is None else str(value)


def find_rows(grid, labels, first_value_col, strict=True):
    """Все строки с дословным путём меток; strict сохранён для старых вызовов."""
    want = [literal(x) for x in labels]
    hits = [r for r in range(len(grid))
            if _match([literal(grid[r][c]) for c in range(first_value_col)], want)]
    return hits, 'exact' if hits else 'none'


def find_column(grid, label, max_scan=10, strict=True):
    """Колонка по заголовку («Month», «Total») — для листов без ряда по месяцам."""
    want = literal(label)
    hits = []
    for r in range(min(max_scan, len(grid))):
        for c, v in enumerate(grid[r]):
            if literal(v) == want:
                hits.append((r, c))
    if len(hits) > 1:
        raise ValueError(f"колонка «{label}» неоднозначна: {len(hits)} совпадений; нужен выбор блока")
    return hits[0] if hits else (None, None)


def read_by_column(grid, ws, item, strict=True, formula_wb=None):
    """Значение на пересечении строки-по-меткам и колонки-по-заголовку.

    Так адресуются оперативные листы: ряда по месяцам нет, есть колонка «Month» и
    строки, различающиеся меткой блока («MACROREGION / Profitability forecast»).
    Координаты в карте не нужны — те же метки, что в описании метрики; строки съедут, а
    метки останутся.
    """
    try:
        hrow, col = find_column(grid, item["column_label"], strict=strict)
    except ValueError as exc:
        return None, str(exc)
    if col is None:
        return None, f"колонка «{item['column_label']}» не найдена"
    hits, mode = find_rows(grid, item["labels"], col, strict=strict)
    if len(hits) != 1:
        return None, ("меток не найдено" if not hits
                      else f"неоднозначно: {len(hits)} строк {[h + 1 for h in hits]}")
    r = hits[0]
    address = ws.cell(r + 1, col + 1).coordinate
    value, err = read_numeric_cell(ws, formula_wb, ws.title, address, item)
    if err:
        return None, f"в строке {r + 1}: {err}"
    return (value, r + 1, mode), None


def read_cell(ws, grid, item, strict=True, formula_wb=None):
    """Адрес одной ячейкой: значение текущего периода, которого нет в ряду по месяцам.

    Так адресуются оперативные показатели («прогноз месяца», «накопленным итогом»):
    ряда по периодам у них нет, есть колонка «Month» и своя строка. Координата без
    проверки опасна — строки в таких листах съезжают, — поэтому карта обязана нести
    guard: ячейку с меткой и её ожидаемый текст. Не сошлось — отказ, не чтение соседа.
    """
    guard = item.get("guard") or {}
    if not guard.get("cell") or not guard.get("label"):
        return None, "адрес ячейкой без guard: нужна метка-подтверждение"
    got = ws[guard["cell"]].value
    if literal(got) != literal(guard["label"]):
        return None, (f"guard не сошёлся: {guard['cell']} = {got!r}, "
                      f"ожидалось {guard['label']!r} — строки съехали")
    return read_numeric_cell(ws, formula_wb, ws.title, item["cell"], item)


def resolve_sheet(wb, name, strict=True):
    """Точное имя листа; strict сохранён для старых вызовов."""
    if name in wb.sheetnames:
        return name, None
    return None, f"лист «{name}» не найден"


def _read_impl(book, card, year=None, upto=None, kind="факт"):
    """upto — последний период, за который бывает ФАКТ (по умолчанию текущий месяц).

    Читается сохранённый числовой результат ячейки; формулы не вычисляются заново.
    Периоды после upto остаются вне факта. План на будущее сохраняется.
    """
    wb = openpyxl.load_workbook(book, data_only=True)
    formulas = None
    try:
        formulas = openpyxl.load_workbook(book, data_only=False)
        return _read_workbook(wb, card, year, upto, kind, formula_wb=formulas)
    finally:
        wb.close()
        if formulas is not None:
            formulas.close()


def _read_workbook(wb, card, year, upto, kind, formula_wb=None):
    cache, out = {}, []
    for item in card["metrics"]:
        strict = True
        row = {"metric": item["name"], "unit": item.get("unit"), "notes": [], "values": {}, "plan": {},
               "companion": bool(item.get("year_offset"))}
        if item.get("note"):
            row["notes"].append(item["note"])
        scale = item.get("scale", 1)
        if scale is None:
            scale = 1
        if isinstance(scale, bool) or not isinstance(scale, (int, float)) or not math.isfinite(scale) or scale <= 0:
            row["status"] = "refused"
            row["reason"] = "масштаб должен быть положительным числом"
            out.append(row)
            continue
        sheet, note = resolve_sheet(wb, item["sheet"], strict=strict)
        if note:
            row["notes"].append(note)
        if sheet is None:
            row["status"] = "refused"; row["reason"] = note; out.append(row); continue
        row["sheet"] = sheet
        if sheet not in cache:
            ws = wb[sheet]; g = sheet_grid(ws)
            try:
                axis, axis_error = find_period_axis(g), None
            except ValueError as exc:
                axis, axis_error = None, str(exc)
            cache[sheet] = (ws, g, axis, axis_error)
        ws, grid, axis, axis_error = cache[sheet]

        if item.get("column_label"):
            got, err = read_by_column(grid, ws, item, strict=strict,
                                      formula_wb=formula_wb)
            if err:
                row["status"] = "refused"; row["reason"] = err; out.append(row); continue
            value, r, mode = got
            row["row"] = r
            period = item.get("period") or upto or "текущий"
            row["values"] = {period: value}
            row["status"] = "ok_with_notes"; out.append(row); continue

        if item.get("cell"):
            value, err = read_cell(ws, grid, item, strict=strict,
                                   formula_wb=formula_wb)
            if err:
                row["status"] = "refused"; row["reason"] = err; out.append(row); continue
            period = item.get("period") or upto or ""
            row["values"] = {period: value} if period else {"текущий": value}
            row["status"] = "ok_with_notes"; out.append(row); continue

        if not axis:
            row["status"] = "refused"; row["reason"] = axis_error or "ось периодов не распознана"; out.append(row); continue
        header_row, cols = axis
        base_year = (year or card.get("year") or 0) + (item.get("year_offset") or 0)
        explicit_years = {int(value) for header in grid[:header_row + 1]
                          for col in cols.values()
                          for value in re.findall(r"\b(?:19|20)\d{2}\b", str(header[col]))}
        if explicit_years and explicit_years != {base_year}:
            row["status"] = "refused"
            row["reason"] = f"год в шапке {sorted(explicit_years)} не совпадает с годом карты {base_year}"
            out.append(row); continue
        first = min(cols.values())

        def series(labels, kind):
            hits, mode = find_rows(grid, labels, first, strict=strict)
            if len(hits) != 1:
                return None, (f"меток не найдено" if not hits
                              else f"неоднозначно: {len(hits)} строк {[h + 1 for h in hits]}"), None
            r = hits[0]
            vals = {}
            for m, c in sorted(cols.items()):
                base_year = (year or card.get('year') or 0) + (item.get("year_offset") or 0)
                period = f"{base_year}-{m:02d}"
                if kind == "факт" and upto and not item.get("year_offset"):
                    if period > upto:
                        continue
                address = ws.cell(r + 1, c + 1).coordinate
                value, err = read_numeric_cell(ws, formula_wb, sheet, address, item)
                if err:
                    row["notes"].append(f"{kind} {period}: {err}")
                    continue
                vals[period] = value
            return r + 1, None, vals

        r, err, vals = series(item["labels"], kind)
        if err:
            row["status"] = "refused"; row["reason"] = err; out.append(row); continue
        row["row"] = r; row["values"] = vals
        if item.get("expected_row") and item["expected_row"] != r:
            row["notes"].append(f"строка была {item['expected_row']}, стала {r} — книга разошлась с описанием метрики")
        if item.get("plan_labels") and "plan_source" not in item:
            pr, perr, pvals = series(item["plan_labels"], "план")
            if perr:
                row["notes"].append(f"план не прочитан: {perr}")
            else:
                row["plan"] = pvals
        row["status"] = "ok" if not row["notes"] else "ok_with_notes"
        out.append(row)
    return out


def read(book, card, year=None, upto=None):
    """Публичный вход: upto по умолчанию — текущий месяц."""
    year = year or card.get("year")
    upto = upto or dt.date.today().strftime("%Y-%m")
    try:
        rows = _read_impl(book, card, year, upto)
    except (OSError, TypeError, ValueError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
        rows = [{"metric": item["name"], "unit": item.get("unit"), "notes": [],
                 "values": {}, "plan": {}, "status": "refused",
                 "reason": f"источник факта не прочитан: {exc}",
                 "companion": bool(item.get("year_offset"))} for item in card["metrics"]]
    for item, row in zip(card["metrics"], rows):
        if "plan_source" not in item:
            continue
        source = item["plan_source"]
        row["plan"] = {}
        try:
            if not isinstance(source, dict):
                raise ValueError("plan_source должен содержать адрес плана")
            plan_item = dict(source)
            plan_item["name"] = item["name"]
            plan_item.setdefault("sheet", item["sheet"])
            plan_item.setdefault("unit", item.get("unit"))
            plan_item.setdefault("scale", 1)
            # Все адреса плана проверяются дословно, как адрес факта.
            plan_item["strict_labels"] = True
            if not plan_item.get("sheet"):
                raise ValueError("лист плана не указан")
            if not plan_item.get("cell") and (not isinstance(plan_item.get("labels"), list)
                                              or not plan_item["labels"]
                                              or any(not str(label).strip() for label in plan_item["labels"])):
                raise ValueError("метки плана не указаны")
            if norm(plan_item["unit"]) != norm(item.get("unit")):
                raise ValueError(f"единица плана «{plan_item['unit']}» отличается от единицы факта «{item.get('unit')}»")
            scale = plan_item["scale"]
            if isinstance(scale, bool) or not isinstance(scale, (int, float)) or not math.isfinite(scale) or scale <= 0:
                raise ValueError("масштаб плана должен быть положительным числом")
            plan_book = source.get("book", book)
            plan_year = source.get("year", year)
            plan_row = _read_impl(plan_book, {"year": plan_year, "metrics": [plan_item]}, plan_year, upto, kind="план")[0]
            if plan_row["status"] == "refused":
                raise ValueError(plan_row["reason"])
            if not plan_row["values"]:
                raise ValueError("числовых значений плана нет")
            row["plan"] = plan_row["values"]
            row["plan_source"] = {"book": str(plan_book), "sheet": plan_row["sheet"], "row": plan_row.get("row")}
            row["notes"].extend(f"план: {note}" for note in plan_row["notes"])
            if source.get("source_note"):
                row["notes"].append(f"план: {source['source_note']}")
        except (OSError, TypeError, ValueError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
            row["notes"].append(f"план не прочитан: {exc}")
            row["plan_status"] = "refused"
        if row["status"] != "refused":
            row["status"] = "ok" if not row["notes"] else "ok_with_notes"
    return rows


def render(rows, source_note=None):
    """Таблица для агента-потребителя: метрика · период · план · факт · пометка."""
    lines = []
    if source_note:
        lines.append(source_note); lines.append("")
    for row in rows:
        head = f"**{row['metric']}**" + (f" ({row['unit']})" if row.get("unit") else "")
        if row["status"] == "refused" and not row["plan"]:
            lines.append(f"{head}: значения нет — {row['reason']}"); lines.append(""); continue
        periods = sorted(set(row["values"]) | set(row["plan"]))
        if row["status"] == "refused":
            head += f" — факт недоступен: {row['reason']}"
        lines.append(head + (f" — {'; '.join(row['notes'])}" if row["notes"] else ""))
        lines.append("| период | " + " | ".join(periods) + " |")
        lines.append("|---" * (len(periods) + 1) + "|")
        lines.append("| факт | " + " | ".join(str(row["values"].get(p, "—")) for p in periods) + " |")
        if row["plan"]:
            lines.append("| план | " + " | ".join(str(row["plan"].get(p, "—")) for p in periods) + " |")
        lines.append("")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--book", required=True, help="снимок книги (.xlsx)")
    ap.add_argument("--card", required=True, help="карта адресов из описаний метрик (.json)")
    ap.add_argument("--upto", help="последний период с фактом, YYYY-MM (по умолчанию текущий месяц)")
    ap.add_argument("--json", action="store_true", help="машинный вывод вместо таблицы")
    a = ap.parse_args()
    card = json.loads(Path(a.card).read_text(encoding="utf-8"))
    upto = a.upto or dt.date.today().strftime("%Y-%m")
    rows = read(a.book, card, card.get("year") or dt.date.today().year, upto)
    if a.json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
    else:
        print(render(rows, card.get("source_note")))
    return 0 if any(r["status"] != "refused" or r["plan"] for r in rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
